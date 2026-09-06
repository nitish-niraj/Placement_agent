"""P5 parser suite — in-memory fixtures covering the P5 exit criteria
(>=2 XLSX, 1 CSV, text PDF, scanned PDF, image; LLM paths mocked)."""

import io
import json

import httpx
import openpyxl
import pytest

from pia_shared.textnorm import digits_only, normalize_name
from pia_worker.ai.provider import NIMProvider
from pia_worker.documents.excel import parse_csv, parse_xlsx
from pia_worker.documents.headers import map_headers
from pia_worker.documents.orchestrator import parse_document_bytes
from pia_worker.documents.pdfdoc import extract_pages, looks_scanned


def _build_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX_VARIANTS = _build_xlsx({
    "Sheet1": [
        ["Sr", "Name of Student", "Roll No.", "Registration No", "Branch"],
        [1, "Aarav Sharma", 2310101, 12515022, "MCA"],
        [2, "priya  singh!", 2310102, 12515087, "MCA"],
        [3, "Rahul Verma", None, "12515101", "MCA"],
        [None, None, None, None, None],  # empty row — skipped
        [4, "Sneha Patel", 2310104, 12515203, "MCA"],
    ],
})
XLSX_MULTI = _build_xlsx({
    "Eligible": [
        ["Enrollment", "Name", "Reg. No"],
        ["EN001", "Dev Mehta", 12516001],
        ["EN002", "Isha Nair", 12516002],
    ],
    "Notes": [["This sheet is not a list"], ["nothing here"]],
})


def test_normalize_name() -> None:
    assert normalize_name("Priya  Singh!") == "priya singh"
    # punctuation becomes separators — both match-sides normalize identically
    assert normalize_name("Dr. A.P.J. Abdul Kalam Jr.") == "a p j abdul kalam"
    assert normalize_name(None) == ""
    assert digits_only("Reg. No- 12515641") == "12515641"


def test_header_mapping_variants() -> None:
    mapping = map_headers(["Sr", "Name of Student", "Roll No.", "Registration No", "Branch"])
    assert mapping[1] == "name"
    assert mapping[2] == "roll_number"
    assert mapping[3] == "registration_number"
    mapping2 = map_headers(["Enrollment", "Name", "Reg. No"])
    assert mapping2[0] == "student_id"
    assert mapping2[2] == "registration_number"


def test_xlsx_header_variants_and_empty_rows() -> None:
    parsed = parse_xlsx(XLSX_VARIANTS)
    assert len(parsed.rows) == 4  # empty row skipped
    assert parsed.parser == "xlsx-v1"
    assert not parsed.needs_review
    first = parsed.rows[0]
    assert first.raw_name == "Aarav Sharma"
    assert first.normalized_name == "aarav sharma"
    assert first.roll_number == "2310101"
    assert first.registration_number == "12515022"
    assert first.other_identifiers == {"Sr": "1", "Branch": "MCA"}
    assert first.source_page_or_sheet == "Sheet1"
    assert parsed.rows[1].normalized_name == "priya singh"
    assert parsed.rows[2].roll_number is None  # missing value stays None — never invented
    assert parsed.rows[2].registration_number == "12515101"


def test_xlsx_multisheet_ignores_non_list_sheet() -> None:
    parsed = parse_xlsx(XLSX_MULTI)
    assert len(parsed.rows) == 2
    assert {r.source_page_or_sheet for r in parsed.rows} == {"Eligible"}
    assert any("Notes" in n for n in parsed.notes)


def test_csv_parsing() -> None:
    csv_bytes = (
        b"Name,Registration No,Roll\n"
        b"Rahul Verma,12515101,2310103\n"
        b"Isha Nair,12516002,\n"
    )
    parsed = parse_csv(csv_bytes)
    assert len(parsed.rows) == 2
    assert parsed.rows[0].registration_number == "12515101"
    assert parsed.rows[1].roll_number is None


def _mock_vision(rows: list[dict]) -> NIMProvider:
    payload = {"rows": rows, "company": "SAFEAEON", "is_candidate_list": True,
               "notes": None}
    return NIMProvider(client=httpx.Client(
        transport=httpx.MockTransport(
            lambda req: httpx.Response(200, json={
                "choices": [{"message": {"content": json.dumps(payload)}}],
                "usage": {}})),
        base_url="http://test"))


NITISH_ROW = [{
    "raw_name": "Nitish Kumar", "normalized_name": "nitish kumar",
    "roll_number": None, "registration_number": "12515641",
    "student_id": None, "other_identifiers": {"branch": "MCA"},
    "extraction_confidence": 0.91, "source_page_or_sheet": "", "row_index": 0,
}]


def test_image_vision_path_mocked() -> None:
    parsed = parse_document_bytes("image/png", "list.png", b"\x89PNG-fake",
                                  provider=_mock_vision(NITISH_ROW), correlation_id="c1")
    assert parsed.parser == "vision-v1"
    assert parsed.is_candidate_list is True
    assert parsed.company == "SAFEAEON"
    assert len(parsed.rows) == 1
    assert parsed.rows[0].registration_number == "12515641"


def test_image_vision_unavailable_falls_back_to_review(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("pia_worker.ai.provider._log_ai_call", lambda *a, **k: None)
    provider = NIMProvider(client=httpx.Client(
        transport=httpx.MockTransport(lambda req: httpx.Response(500, json={})),
        base_url="http://test"))

    from pia_worker.documents.types import ParsedDocument

    def fake_ocr(data: bytes) -> ParsedDocument:
        return ParsedDocument(text="NAME REG\nNITISH 12515641", parser="rapidocr-v1",
                              confidence=0.4, needs_review=True,
                              notes=["rapidocr fallback"])

    monkeypatch.setattr("pia_worker.documents.images.rapid_ocr", fake_ocr)
    parsed = parse_document_bytes("image/png", "list.png", b"\x89PNG-fake",
                                  provider=provider, correlation_id="c2")
    assert parsed.parser == "rapidocr-v1"
    assert parsed.needs_review is True


def test_pdf_text_vs_scanned() -> None:
    import pymupdf

    doc = pymupdf.Document()
    page = doc.new_page()
    page.insert_text((72, 72), "Name Reg No\nNitish Kumar 12515641\nAarav Sharma 12515022")
    pdf_text = doc.tobytes()
    pages = extract_pages(pdf_text)
    assert not looks_scanned(pages)
    assert "Nitish Kumar" in pages[0]["text"]

    doc2 = pymupdf.Document()
    page2 = doc2.new_page(width=200, height=100)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100))
    page2.insert_image(page2.rect, pixmap=pix)
    assert looks_scanned(extract_pages(doc2.tobytes()))


def test_unsupported_type_needs_review() -> None:
    parsed = parse_document_bytes("application/zip", "thing.zip", b"PK", provider=None)
    assert parsed.needs_review
    assert parsed.rows == []


def test_low_confidence_rows_flag_review() -> None:
    weak_rows = [{**NITISH_ROW[0], "raw_name": "smudged name",
                  "normalized_name": "smudged name", "registration_number": None,
                  "extraction_confidence": 0.4}]
    parsed = parse_document_bytes("image/png", "blur.png", b"\x89PNG",
                                  provider=_mock_vision(weak_rows))
    assert parsed.needs_review  # below vision_confidence_floor (0.75)
    assert any("below confidence floor" in n for n in parsed.notes)


def test_candidate_row_schema_never_invents() -> None:
    """Absent identifiers must stay None (master §8.1)."""
    from pia_shared.schemas import CandidateRow

    row = CandidateRow(raw_name="X", normalized_name="x")
    assert row.roll_number is None and row.registration_number is None
    with pytest.raises(ValueError):
        CandidateRow(raw_name="X", normalized_name="x", extraction_confidence=5)
